"""Post-allocation operations: booking, three-way reconciliation, surveillance, KPIs/KRIs, Excel report.

Preventive controls (eligibility, sizing, co-investment checks, maker-checker) stop bad allocations
before booking. The controls here are DETECTIVE: they independently re-check what actually happened
downstream – what was booked, what the ledger shows, and where each fund's portfolio now stands.
BDC figures and thresholds are illustrative and simplified.
"""
from __future__ import annotations

import json
import sqlite3
from datetime import datetime, timezone

from . import audit

TOL = 1e-6
QUALIFYING_MIN = 0.70        # BDC: qualifying assets >= 70% of total assets (illustrative, simplified)
ASSET_COVERAGE_MIN = 1.50    # BDC: asset coverage >= 150% (illustrative, simplified)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _next(con, table, col, prefix):
    n = con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] + 1
    while con.execute(f"SELECT 1 FROM {table} WHERE {col}=?", (f"{prefix}{n:04d}",)).fetchone():
        n += 1
    return f"{prefix}{n:04d}"


def _exception(con, run_id, fund_id, rule_id, severity, owner, detail, escalation):
    from .allocate import _due
    exc_id, opened = _next(con, "exceptions", "exc_id", "E"), _now()
    con.execute("INSERT INTO exceptions (exc_id, run_id, fund_id, rule_id, severity, status, owner, opened_at, detail, due_at) "
                "VALUES (?,?,?,?,?,'Open',?,?,?,?)",
                (exc_id, run_id, fund_id, rule_id, severity, owner, opened, detail, _due(opened, severity, escalation or {})))
    return exc_id


# ---------------------------------------------------------------- booking
def parse_break(spec: str | None) -> dict:
    """'F3:-0.5' -> {'F3': -0.5};  'F3:missing' -> {'F3': 'missing'} (test injection only)."""
    if not spec:
        return {}
    fund, val = spec.split(":", 1)
    return {fund: "missing" if val == "missing" else float(val)}


def book(con: sqlite3.Connection, run_id: str, by: str, inject_break: dict | None = None) -> dict:
    r = con.execute("SELECT deal_id, mode, status FROM allocation_runs WHERE run_id=?", (run_id,)).fetchone()
    if r is None:
        raise ValueError(f"unknown run {run_id}")
    deal_id, mode, status = r
    if mode != "live" or status != "approved":
        raise ValueError(f"only approved live runs can be booked ({run_id} is {mode}/{status})")
    other = con.execute("SELECT run_id FROM allocation_runs WHERE deal_id=? AND status IN ('booked','reconciled') "
                        "AND run_id<>?", (deal_id, run_id)).fetchone()
    if other:
        raise ValueError(f"deal {deal_id} already booked under {other[0]} – double booking refused")
    issuer, qualifying = con.execute("SELECT issuer, qualifying FROM deals WHERE deal_id=?", (deal_id,)).fetchone()
    from . import bdc
    lines = con.execute("SELECT fund_id, final_amt FROM allocation_lines WHERE run_id=? AND final_amt>0 "
                        "ORDER BY fund_id", (run_id,)).fetchall()
    funding = {}
    for fid, _ in lines:                                           # B07: a BDC line cannot be booked on unknown funding
        if bdc.portfolio(con, fid) is not None:
            f = bdc.resolve_funding(con, deal_id, fid)
            if f["source"] == "unknown":
                raise ValueError(f"funding for {fid} on {deal_id} is not confirmed – run set-funding before booking")
            funding[fid] = f
    inject_break = inject_break or {}
    ts, booked = _now(), {}
    for fid, amt in lines:
        con.execute("INSERT INTO bookings VALUES (?,?,?,?)", (run_id, fid, amt, ts))
        inj = inject_break.get(fid)
        if inj != "missing":                                         # downstream ledger posting
            con.execute("INSERT INTO ledger_stub VALUES (?,?,?,?)", (fid, deal_id, round(amt + (inj or 0.0), 9), ts))
        con.execute("INSERT INTO holdings VALUES (?,?,?) ON CONFLICT(fund_id, issuer) DO UPDATE "
                    "SET exposure = exposure + excluded.exposure", (fid, issuer, amt))
        con.execute("UPDATE funds SET available_capital = available_capital - ? WHERE fund_id=?", (amt, fid))
        if fid in funding:                                         # BDC balance sheet moves with the funding (B07)
            st = bdc.apply_purchase(bdc.portfolio(con, fid), amt, funding[fid]["cash_share"], bool(qualifying))
            prm = json.loads(con.execute("SELECT params FROM rules WHERE rule_id='SIZE-BDC-QUALIFYING-HEADROOM' "
                                         "ORDER BY version DESC").fetchone()[0])
            if (not qualifying and st["Q"] / st["T"] < prm["min_ratio"] - 1e-9) or \
               (st["S"] and st["T"] / st["S"] < prm.get("min_asset_coverage", 1.5) - 1e-9) or st["C"] < -1e-9:
                con.rollback()
                raise ValueError(f"pre-booking re-check failed for {fid}: qualifying {st['Q']/st['T']:.3%}, "
                                 f"coverage {st['T']/st['S']:.1%}, cash {st['C']:g} – escalate to Compliance")
            con.execute("UPDATE fund_portfolio SET qualifying_assets=?, total_assets=?, senior_securities=?, cash=? "
                        "WHERE fund_id=?", (st["Q"], st["T"], st["S"], st["C"], fid))
        booked[fid] = amt
    con.execute("UPDATE allocation_runs SET status='booked' WHERE run_id=?", (run_id,))
    audit.append(con, by, "BOOKED", run_id, {"deal_id": deal_id, "booked": booked,
                                             "funding": {k: v["source"] for k, v in funding.items()},
                                             "test_injected_ledger_break": inject_break or None})
    con.commit()
    return booked


# ---------------------------------------------------------------- reconciliation
def reconcile(con: sqlite3.Connection, run_id: str, by: str = "ops1", escalation: dict | None = None) -> dict:
    r = con.execute("SELECT deal_id, status, unallocated FROM allocation_runs WHERE run_id=?", (run_id,)).fetchone()
    if r is None:
        raise ValueError(f"unknown run {run_id}")
    deal_id, status, unalloc = r
    if status not in ("booked", "reconciled"):
        raise ValueError(f"run {run_id} must be booked before reconciliation (is '{status}')")
    approved = dict(con.execute("SELECT fund_id, final_amt FROM allocation_lines WHERE run_id=? AND final_amt>0", (run_id,)))
    booked = dict(con.execute("SELECT fund_id, amount FROM bookings WHERE run_id=?", (run_id,)))
    ledger = dict(con.execute("SELECT fund_id, SUM(amount) FROM ledger_stub WHERE deal_id=? GROUP BY fund_id", (deal_id,)))
    size = con.execute("SELECT size FROM deals WHERE deal_id=?", (deal_id,)).fetchone()[0]
    breaks = []

    def brk(fund, check, expected, actual):
        bid = _next(con, "recon_breaks", "break_id", "B")
        diff = None if expected is None or actual is None else round(actual - expected, 9)
        con.execute("INSERT INTO recon_breaks (break_id, run_id, fund_id, check_name, expected, actual, diff, status, "
                    "opened_at) VALUES (?,?,?,?,?,?,?,'Open',?)",
                    (bid, run_id, fund, check, expected, actual, diff, _now()))
        exc = _exception(con, run_id, fund, f"RECON-{check}", "High", "Operations",
                         f"{check} break for {fund}: expected {expected}, actual {actual}", escalation)
        breaks.append({"break_id": bid, "fund_id": fund, "check": check, "expected": expected,
                       "actual": actual, "diff": diff, "exc_id": exc})

    for f in sorted(set(approved) | set(booked)):
        a, b = approved.get(f), booked.get(f)
        if a is None or b is None or abs(a - b) > TOL:
            brk(f, "APPROVED_vs_BOOKED", a, b)
    for f in sorted(set(booked) | set(ledger)):
        b, l = booked.get(f), ledger.get(f)
        if b is None or l is None or abs(b - l) > TOL:
            brk(f, "BOOKED_vs_LEDGER", b, l)
    if abs(sum(approved.values()) + unalloc - size) > TOL:
        brk(None, "CONTROL_TOTAL", size, round(sum(approved.values()) + unalloc, 9))

    con.execute("UPDATE allocation_runs SET status='reconciled' WHERE run_id=?", (run_id,))
    summary = {"run_id": run_id, "funds_checked": len(set(approved) | set(booked) | set(ledger)),
               "checks": ["APPROVED_vs_BOOKED", "BOOKED_vs_LEDGER", "CONTROL_TOTAL"], "breaks": breaks}
    audit.append(con, by, "RECONCILED", run_id, {"breaks": [(b["break_id"], b["fund_id"], b["check"], b["diff"]) for b in breaks]})
    con.commit()
    return summary


def resolve_break(con: sqlite3.Connection, break_id: str, by: str, explanation: str,
                  post_adjustment: bool = True) -> dict:
    """Close a reconciliation break with a root cause. For a BOOKED_vs_LEDGER break, optionally post an adjusting
    ledger entry, then re-check that the fund now ties out before marking the break and its exception Resolved."""
    if not (explanation and explanation.strip()):
        raise ValueError("a break can only be resolved with an explanation (root cause and action)")
    r = con.execute("SELECT run_id, fund_id, check_name, expected, actual, diff, status FROM recon_breaks "
                    "WHERE break_id=?", (break_id,)).fetchone()
    if r is None:
        raise ValueError(f"unknown break {break_id}")
    run_id, fund, check, expected, actual, diff, status = r
    if status != "Open":
        raise ValueError(f"break {break_id} is already {status}")
    deal_id = con.execute("SELECT deal_id FROM allocation_runs WHERE run_id=?", (run_id,)).fetchone()[0]
    adjustment = None
    if check == "BOOKED_vs_LEDGER" and post_adjustment:
        adjustment = round(expected - (actual or 0.0), 9)
        con.execute("INSERT INTO ledger_stub VALUES (?,?,?,?)", (fund, deal_id, adjustment, _now()))
        ledger = con.execute("SELECT SUM(amount) FROM ledger_stub WHERE fund_id=? AND deal_id=?", (fund, deal_id)).fetchone()[0]
        if abs(ledger - expected) > TOL:
            raise ValueError(f"adjustment did not clear the break: ledger {ledger} vs booked {expected}")
    resolution = explanation + (f" | adjusting ledger entry {adjustment:+g} posted; fund re-checked and ties out"
                                if adjustment is not None else "")
    ts = _now()
    con.execute("UPDATE recon_breaks SET status='Resolved', resolved_at=?, resolution=? WHERE break_id=?",
                (ts, resolution, break_id))
    exc = [e[0] for e in con.execute("SELECT exc_id FROM exceptions WHERE run_id=? AND rule_id=? AND fund_id IS ? "
                                     "AND status='Open'", (run_id, f"RECON-{check}", fund))]
    for e in exc:
        con.execute("UPDATE exceptions SET status='Resolved', resolved_at=?, resolution=? WHERE exc_id=?",
                    (ts, f"{break_id} resolved by {by}: {resolution}", e))
    audit.append(con, by, "BREAK_RESOLVED", run_id, {"break_id": break_id, "fund_id": fund, "check": check,
                                                      "adjustment": adjustment, "explanation": explanation,
                                                      "resolved_exceptions": exc})
    con.commit()
    return {"break_id": break_id, "adjustment": adjustment, "resolved_exceptions": exc}


# ---------------------------------------------------------------- surveillance
def surveil(con: sqlite3.Connection, by: str = "surveillance", escalation: dict | None = None) -> list[dict]:
    ts, results = _now(), []
    funds = {r[0]: {"nav": r[1], "cap": r[2], "type": r[3]} for r in
             con.execute("SELECT fund_id, nav, max_issuer_pct, vehicle_type FROM funds")}

    def record(fund, metric, value, threshold, breach, detail):
        cid = _next(con, "surveillance_results", "check_id", "S")
        con.execute("INSERT INTO surveillance_results VALUES (?,?,?,?,?,?,?,?)",
                    (cid, ts, fund, metric, value, threshold, int(breach), detail))
        res = {"check_id": cid, "fund_id": fund, "metric": metric, "value": value, "threshold": threshold,
               "breach": breach, "detail": detail}
        if breach:
            res["exc_id"] = _exception(con, None, fund, f"SURV-{metric.split(':')[0]}", "High", "Compliance",
                                       detail, escalation)
        results.append(res)

    for fid, issuer, exposure in con.execute("SELECT fund_id, issuer, exposure FROM holdings ORDER BY fund_id, issuer").fetchall():
        f = funds[fid]
        pct = round(exposure / f["nav"], 6)
        record(fid, f"ISSUER_CONCENTRATION:{issuer}", pct, f["cap"], pct > f["cap"] + TOL,
               f"{issuer} = {pct:.2%} of NAV (limit {f['cap']:.0%})")
    for fid, total, qual, senior in con.execute("SELECT fund_id, total_assets, qualifying_assets, senior_securities "
                                                "FROM fund_portfolio ORDER BY fund_id").fetchall():
        q = round(qual / total, 6)
        record(fid, "QUALIFYING_ASSETS_RATIO", q, QUALIFYING_MIN, q < QUALIFYING_MIN - TOL,
               f"qualifying assets {q:.2%} of total (minimum {QUALIFYING_MIN:.0%}; illustrative)")
        cov = round(total / senior, 6)
        record(fid, "ASSET_COVERAGE", cov, ASSET_COVERAGE_MIN, cov < ASSET_COVERAGE_MIN - TOL,
               f"asset coverage {cov:.0%} (minimum {ASSET_COVERAGE_MIN:.0%}; illustrative, simplified)")
    audit.append(con, by, "SURVEILLANCE_RUN", "portfolio",
                 {"checks": len(results), "breaches": [(r["fund_id"], r["metric"]) for r in results if r["breach"]]})
    con.commit()
    return results


# ---------------------------------------------------------------- KPIs / KRIs
def metrics(con: sqlite3.Connection) -> dict:
    q = lambda sql, *a: con.execute(sql, a).fetchone()[0]
    runs_total = q("SELECT COUNT(*) FROM allocation_runs")
    by_status = dict(con.execute("SELECT status, COUNT(*) FROM allocation_runs GROUP BY status"))
    with_exc = q("SELECT COUNT(DISTINCT run_id) FROM exceptions WHERE run_id IS NOT NULL")
    created = {e: datetime.fromisoformat(t) for e, t in
               con.execute("SELECT entity, ts FROM audit_log WHERE action='RUN_CREATED'")}
    cycles = [(datetime.fromisoformat(t) - created[e]).total_seconds() / 60
              for e, t in con.execute("SELECT entity, ts FROM audit_log WHERE action='APPROVED'") if e in created]
    now = _now()
    return {
        "KPI": {
            "runs_total": runs_total,
            "runs_by_status": by_status,
            "live_runs_approved_or_later": q("SELECT COUNT(*) FROM allocation_runs WHERE mode='live' "
                                             "AND status IN ('approved','booked','reconciled')"),
            "avg_create_to_approval_minutes": round(sum(cycles) / len(cycles), 2) if cycles else None,
            "unallocated_usd_m_on_live_runs": q("SELECT COALESCE(SUM(unallocated),0) FROM allocation_runs WHERE mode='live'"),
            "exceptions_resolved": q("SELECT COUNT(*) FROM exceptions WHERE status IN ('Resolved','Superseded')"),
        },
        "KRI": {
            "runs_with_exceptions_pct": round(100 * with_exc / runs_total, 1) if runs_total else 0.0,
            "open_exceptions_by_severity": dict(con.execute(
                "SELECT severity, COUNT(*) FROM exceptions WHERE status='Open' GROUP BY severity")),
            "overdue_open_exceptions": q("SELECT COUNT(*) FROM exceptions WHERE status='Open' AND due_at IS NOT NULL "
                                         "AND due_at < ?", now),
            "approvals_with_override_rationale": q("SELECT COUNT(*) FROM approvals WHERE decision='approve' "
                                                   "AND rationale IS NOT NULL AND rationale<>''"),
            "open_recon_breaks": q("SELECT COUNT(*) FROM recon_breaks WHERE status='Open'"),
            "surveillance_breaches_latest": q("SELECT COUNT(*) FROM surveillance_results WHERE breach=1 AND ts="
                                              "(SELECT MAX(ts) FROM surveillance_results)") if q(
                "SELECT COUNT(*) FROM surveillance_results") else 0,
            "audit_chain_ok": audit.verify(con)[0],
        },
    }


# ---------------------------------------------------------------- Excel report
def export_excel(con: sqlite3.Connection, path: str) -> list[str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, sql, money_cols=(), pct_cols=()):
        ws = wb.create_sheet(title)
        cur = con.execute(sql)
        cols = [c[0] for c in cur.description]
        ws.append(cols)
        for row in cur.fetchall():
            ws.append(list(row))
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
        for i, name in enumerate(cols, start=1):
            letter = get_column_letter(i)
            width = max(len(str(name)), *(len(str(v.value)) for v in ws[letter][1:]), 8) if ws.max_row > 1 else len(name) + 2
            ws.column_dimensions[letter].width = min(width + 2, 60)
            fmt = "#,##0.0" if name in money_cols else "0.00%" if name in pct_cols else None
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    m = metrics(con)
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Allocation engine – management report (synthetic data, illustrative policies)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Generated {_now()}"])
    ws.append([])
    for group in ("KPI", "KRI"):
        ws.append([group, "Value"])
        for c in ws[ws.max_row]:
            c.font, c.fill = head_font, head_fill
        for k, v in m[group].items():
            ws.append([k, json.dumps(v) if isinstance(v, dict) else v])
        ws.append([])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 45

    sheet("Allocations", "SELECT r.run_id, r.deal_id, d.issuer, r.mode, r.status, r.allocation_key, l.fund_id, "
          "f.vehicle_type, l.final_amt, l.sizing_notes FROM allocation_lines l JOIN allocation_runs r USING(run_id) "
          "JOIN deals d ON d.deal_id=r.deal_id JOIN funds f ON f.fund_id=l.fund_id "
          "WHERE r.mode='live' AND l.final_amt>0 ORDER BY r.run_id, l.fund_id", money_cols=("final_amt",))
    # Fairness (B06): per-fund fair share vs actual; drift and drift % are live Excel formulas
    from .reports import fairness
    fr = fairness(con)
    wf = wb.create_sheet("Fairness")
    wf.append(["Fund", "Deals eligible", "Deals excluded", "Fair share", "Actual", "Drift", "Drift %", "Flag (|drift %| > tolerance)",
               "Drift reasons", "Exclusion reasons"])
    for c in wf[1]:
        c.font, c.fill = head_font, head_fill
    wf["L1"], wf["M1"] = "Tolerance", fr["tolerance"]
    wf["L1"].font = Font(bold=True)
    for i, f in enumerate(fr["funds"], start=2):
        wf.append([f["fund_id"], f["deals_eligible"], f["deals_excluded"], f["fair"], f["actual"], f"=E{i}-D{i}",
                   f"=IF(D{i}=0,0,F{i}/D{i})", f'=IF(ABS(G{i})>$M$1,"REVIEW","OK")',
                   ", ".join(f"{k} x{v}" for k, v in sorted(f["reasons"].items())),
                   ", ".join(f"{k} x{v}" for k, v in sorted(f["exclusions"].items()))])
    n = len(fr["funds"]) + 1
    wf.append(["Total", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=E{n+1}-D{n+1}"])
    for c in wf[n + 1]:
        c.font = Font(bold=True)
    for row in wf.iter_rows(min_row=2, max_row=n + 1):
        for c in row[3:6]:
            c.number_format = "#,##0.00"
        row[6].number_format = "0.00%"
    wf["M1"].number_format = "0%"
    wf.append([])
    wf.append(["Method: " + fr["method"]])
    for col, w in zip("ABCDEFGHIJ", (8, 14, 14, 12, 12, 10, 10, 26, 60, 50)):
        wf.column_dimensions[col].width = w
    wf.freeze_panes = "A2"
    sheet("Fairness by deal", "SELECT 1 AS placeholder WHERE 0")  # replaced below
    wd = wb["Fairness by deal"]
    wd.delete_rows(1, wd.max_row)
    wd.append(["Run", "Deal", "Key", "Fund", "Fair share", "Actual", "Drift", "Reason codes"])
    for c in wd[1]:
        c.font, c.fill = head_font, head_fill
    for i, d in enumerate(fr["deals"], start=2):
        wd.append([d["run_id"], d["deal_id"], d["key"], d["fund_id"], d["fair"], d["actual"], f"=F{i}-E{i}", d["reasons"]])
        for c in wd[i][4:7]:
            c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFGH", (8, 8, 30, 6, 12, 12, 10, 70)):
        wd.column_dimensions[col].width = w
    wd.freeze_panes = "A2"
    sheet("Exceptions", "SELECT exc_id, run_id, fund_id, rule_id, severity, status, owner, opened_at, due_at, "
          "resolution, detail FROM exceptions ORDER BY exc_id")
    sheet("Reconciliation", "SELECT break_id, run_id, fund_id, check_name, expected, actual, diff, status, opened_at "
          "FROM recon_breaks ORDER BY break_id", money_cols=("expected", "actual", "diff"))
    sheet("Surveillance", "SELECT check_id, ts, fund_id, metric, value, threshold, breach, detail FROM surveillance_results "
          "ORDER BY check_id")
    sheet("Approvals", "SELECT run_id, maker, checker, decision, rationale, ts FROM approvals ORDER BY ts")
    sheet("Audit trail", "SELECT seq, ts, actor, action, entity, payload, hash FROM audit_log ORDER BY seq")
    wb.save(path)
    return wb.sheetnames
